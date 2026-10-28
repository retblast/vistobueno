"""Carátula de tesis con la **estructura real** (auditoría 2026-10-26).

Este módulo construye el DOCX mínimo que reproduce la estructura real de
la carátula de una tesis de la UNT:

- la etiqueta `AUTORES:` va en su **propio párrafo** y cada autor en un
  párrafo aparte, con la línea de ORCID intercalada;
- la etiqueta `ASESORA:` va en su propio párrafo y el nombre en el
  siguiente;
- el título termina en ", Trujillo <año>" y la portada cierra con la
  línea de ciudad y el año en párrafos separados;
- las palabras clave están en el RESUMEN, no en la carátula.

Los nombres son **placeholders**: ningún dato real de una tesis entra al
repositorio.

Uso:
    from _caratula_real import capturar_ficha, crear_docx_caratula
"""

import tempfile
from pathlib import Path

from docx import Document

from validator.engine import build_report, load_ficha_rules, load_rules, validate_docx

RAIZ = Path(__file__).resolve().parent.parent
REGLAS_YAML = RAIZ / "reglas_unt.yaml"
REGLAS_FICHA_YAML = RAIZ / "reglas_ficha.yaml"

# --- Contenido de la carátula (placeholders, estructura real) -------------

PORTADA = [
    "UNIVERSIDAD NACIONAL DE TRUJILLO",
    "FACULTAD DE EDUCACIÓN Y CIENCIAS DE LA COMUNICACIÓN",
    "ESCUELA PROFESIONAL DE EDUCACIÓN PRIMARIA",
    "Título del trabajo de investigación: Estrategias lúdicas para mejorar "
    "la identidad cultural en estudiantes de primaria de una institución "
    "educativa, Trujillo 2024",
    "TESIS",
    "Para Optar el Título Profesional de Licenciado en Educación Primaria",
    "AUTORES:",
    "Br. Apellido, Nombre",
    "https://orcid.org/0000-0000-0000-0000",
    "Br. Apellido2, Nombre2",
    "https://orcid.org/0000-0000-0000-0001",
    "ASESORA:",
    "Dra. Apellido, Nombre",
    "https://orcid.org/0000-0000-0000-0002",
    "Línea de Investigación:",
    "Educación y Ciencias de la Comunicación y Desarrollo Sostenible",
    "Trujillo – Perú",
    "2026",
]

# --- Resumen: las palabras clave viven aquí, no en la carátula ------------

RESUMEN = [
    "RESUMEN",
    "Palabras Clave: Educación, Cultura, Identidad cultural.",
]


def crear_docx_caratula() -> str:
    """Escribe el DOCX de la carátula real y devuelve su ruta."""
    documento = Document()
    for linea in PORTADA + RESUMEN:
        documento.add_paragraph(linea)
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        ruta = tmp.name
    documento.save(ruta)
    return ruta


def capturar_ficha(ruta: str) -> dict:
    """Corre el motor y devuelve `reporte["ficha"]` de ese DOCX."""
    resultados = validate_docx(
        ruta,
        load_rules(str(REGLAS_YAML)),
        load_ficha_rules(str(REGLAS_FICHA_YAML)),
    )
    return build_report(resultados).get("ficha", {})


def valor(ficha: dict, campo: str):
    """El valor capturado de un campo de la ficha."""
    return (ficha.get(campo) or {}).get("valor")
